"""
CogniHire — Document Parser Module
Extracts raw text from PDF and DOCX resume files.
"""
import io
import re
from typing import Optional


def parse_pdf(file_bytes: bytes) -> Optional[str]:
    """Extract text from a PDF file."""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts) if text_parts else None
    except Exception as e:
        print(f"PDF parsing error: {e}")
        return None


def parse_docx(file_bytes: bytes) -> Optional[str]:
    """Extract text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        return "\n".join(text_parts) if text_parts else None
    except Exception as e:
        print(f"DOCX parsing error: {e}")
        return None


def parse_resume(file_bytes: bytes, filename: str) -> Optional[str]:
    """Parse a resume file based on its extension."""
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext in ("docx", "doc"):
        return parse_docx(file_bytes)
    else:
        return None


def clean_text(text: str) -> str:
    """Clean and normalize extracted text."""
    if not text:
        return ""
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = "".join(char for char in text if char.isprintable() or char in "\n\t")
    return text.strip()
